import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any

THIN_BORDER = Border(
    left=Side(style='thin', color='D0D5DD'),
    right=Side(style='thin', color='D0D5DD'),
    top=Side(style='thin', color='D0D5DD'),
    bottom=Side(style='thin', color='D0D5DD')
)

HEADER_FILL = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')

FONT_BOLD_14 = Font(name='Calibri', size=14, bold=True)
FONT_BOLD_12 = Font(name='Calibri', size=12, bold=True)
FONT_BOLD_10 = Font(name='Calibri', size=10, bold=True)
FONT_REGULAR_10 = Font(name='Calibri', size=10)
FONT_BOLD_9 = Font(name='Calibri', size=9, bold=True)
FONT_REGULAR_9 = Font(name='Calibri', size=9)

def build_color_map(stamps: List[Dict[str, Any]]) -> Dict[str, PatternFill]:
    color_map = {}
    if not stamps:
        return color_map
    for st in stamps:
        code = str(st.get("code", "")).strip().upper()
        bg = st.get("bg_color", "#F1F5F9").lstrip("#")
        if len(bg) == 6:
            color_map[code] = PatternFill(start_color=bg, end_color=bg, fill_type='solid')
    return color_map

def generate_company_excel(
    file_path: str,
    property_name: str,
    property_full: str,
    company_name: str,
    department: str,
    month_info: Dict[str, Any],
    associates: List[Dict[str, Any]],
    settings: Dict[str, Any] = None,
    stamps: List[Dict[str, Any]] = None,
    authorizers: List[Dict[str, Any]] = None
):
    wb = Workbook()
    color_map = build_color_map(stamps or [])
    
    stamp_hours = {}
    if stamps:
        for st in stamps:
            c = str(st.get("code", "")).strip().upper()
            stamp_hours[c] = float(st.get("hours", 0))
    else:
        stamp_hours = {"8": 8, "WO": 8, "PH": 8, "SL": 8, "OFF": 8}
        
    paid_codes = [c for c, h in stamp_hours.items() if h > 0]
    
    # SHEET 1: Daily Attendance
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.views.sheetView[0].showGridLines = True
    
    days_cnt = month_info["days_in_month"]
    period_str = month_info["period_str"]
    
    ws1["A2"] = property_full
    ws1["A2"].font = FONT_BOLD_14
    
    ws1["A3"] = "Daily Attendance Transaction Report"
    ws1["A3"].font = FONT_BOLD_12
    
    ws1["A4"] = f" Transaction Period From  {period_str}"
    ws1["A4"].font = FONT_REGULAR_10
    
    ws1["C5"] = company_name.upper()
    ws1["C5"].font = FONT_BOLD_12
    
    ws1["C6"] = f"Department: {department}"
    ws1["C6"].font = FONT_REGULAR_10
    
    ws1["B8"] = "SL"
    ws1["B8"].font = FONT_BOLD_10
    ws1["B8"].alignment = Alignment(horizontal='center', vertical='center')
    ws1["B8"].border = THIN_BORDER
    ws1["B8"].fill = HEADER_FILL
    
    ws1["C8"] = "Employee Name"
    ws1["C8"].font = FONT_BOLD_10
    ws1["C8"].alignment = Alignment(horizontal='left', vertical='center')
    ws1["C8"].border = THIN_BORDER
    ws1["C8"].fill = HEADER_FILL
    
    ws1["E8"] = "Days"
    ws1["E8"].font = FONT_BOLD_10
    ws1["E8"].alignment = Alignment(horizontal='center', vertical='center')
    ws1["E8"].fill = HEADER_FILL
    
    for d in range(1, days_cnt + 1):
        col_letter = get_column_letter(4 + d)
        cell = ws1[f"{col_letter}9"]
        cell.value = d
        cell.font = FONT_BOLD_9
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER
        cell.fill = HEADER_FILL
        
    tot_hrs_col = get_column_letter(5 + days_cnt)
    work_col = get_column_letter(6 + days_cnt)
    
    ws1[f"{tot_hrs_col}9"] = "Total hours"
    ws1[f"{tot_hrs_col}9"].font = FONT_BOLD_9
    ws1[f"{tot_hrs_col}9"].alignment = Alignment(horizontal='center', vertical='center')
    ws1[f"{tot_hrs_col}9"].border = THIN_BORDER
    ws1[f"{tot_hrs_col}9"].fill = HEADER_FILL
    
    ws1[f"{work_col}9"] = "Work"
    ws1[f"{work_col}9"].font = FONT_BOLD_9
    ws1[f"{work_col}9"].alignment = Alignment(horizontal='center', vertical='center')
    ws1[f"{work_col}9"].border = THIN_BORDER
    ws1[f"{work_col}9"].fill = HEADER_FILL
    
    current_row = 10
    first_d_col = get_column_letter(5)
    last_d_col = get_column_letter(4 + days_cnt)
    
    for idx, assoc in enumerate(associates, start=1):
        c_sl = ws1[f"B{current_row}"]
        c_sl.value = idx
        c_sl.font = FONT_REGULAR_10
        c_sl.alignment = Alignment(horizontal='center', vertical='center')
        c_sl.border = THIN_BORDER
        
        c_name = ws1[f"C{current_row}"]
        c_name.value = assoc["name"]
        c_name.font = FONT_REGULAR_10
        c_name.alignment = Alignment(horizontal='left', vertical='center')
        c_name.border = THIN_BORDER
        
        days_vals = assoc.get("days", ["8"] * days_cnt)
        
        for d_idx in range(days_cnt):
            raw_val = days_vals[d_idx] if d_idx < len(days_vals) else "8"
            code = str(raw_val).strip().upper()
            c_d = ws1.cell(row=current_row, column=5 + d_idx)
            
            if code in ["8", 8]:
                c_d.value = 8
                c_d.font = FONT_REGULAR_9
            else:
                c_d.value = code
                c_d.font = FONT_BOLD_9
                if code in color_map:
                    c_d.fill = color_map[code]
                    
            c_d.alignment = Alignment(horizontal='center', vertical='center')
            c_d.border = THIN_BORDER
            
        c_tot = ws1[f"{tot_hrs_col}{current_row}"]
        c_work = ws1[f"{work_col}{current_row}"]
        
        work_parts = []
        if "8" in paid_codes:
            work_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, 8)")
            work_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, \"8\")")
        for pc in paid_codes:
            if pc != "8":
                work_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, \"{pc}\")")
                
        c_work.value = "=" + " + ".join(work_parts) if work_parts else "0"
        c_work.font = FONT_BOLD_10
        c_work.alignment = Alignment(horizontal='center', vertical='center')
        c_work.border = THIN_BORDER
        
        unique_hours = set(stamp_hours[c] for c in paid_codes)
        if len(unique_hours) == 1 and 8 in unique_hours:
            c_tot.value = f"={work_col}{current_row}*8"
        else:
            tot_parts = []
            if "8" in paid_codes:
                h = stamp_hours["8"]
                tot_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, 8)*{int(h) if h.is_integer() else h}")
                tot_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, \"8\")*{int(h) if h.is_integer() else h}")
            for pc in paid_codes:
                if pc != "8":
                    h = stamp_hours[pc]
                    tot_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, \"{pc}\")*{int(h) if h.is_integer() else h}")
            c_tot.value = "=" + " + ".join(tot_parts) if tot_parts else "0"
            
        c_tot.font = FONT_BOLD_10
        c_tot.alignment = Alignment(horizontal='center', vertical='center')
        c_tot.border = THIN_BORDER
        
        current_row += 1
        
    # Dynamic Signatures Row: exactly 6 boxes (rows) after the grid table
    sig_row = current_row + 6
    if authorizers is not None:
        active_auths = [a for a in authorizers if a.get("is_active", True)]
    else:
        active_auths = [
            {"title": "Department Head", "name": ""},
            {"title": "Director of Human Resource", "name": ""},
            {"title": "Director of Finance", "name": ""}
        ]
    if not active_auths:
        active_auths = [{"title": "Department Head", "name": ""}]
        
    n_auths = len(active_auths)
    start_c = 4  # Column D
    end_c = 4 + days_cnt  # Near end column
    
    for i, auth in enumerate(active_auths):
        if n_auths == 1:
            col_idx = (start_c + end_c) // 2
        else:
            col_idx = int(start_c + i * ((end_c - start_c) / (n_auths - 1)))
        col_letter = get_column_letter(col_idx)
        cell = ws1[f"{col_letter}{sig_row}"]
        cell.value = auth.get("title", "Authorizer")
        cell.font = FONT_BOLD_10
        cell.alignment = Alignment(horizontal='center')
        if auth.get("name", "").strip():
            c_sub = ws1[f"{col_letter}{sig_row + 1}"]
            c_sub.value = f"({auth['name'].strip()})"
            c_sub.font = FONT_REGULAR_9
            c_sub.alignment = Alignment(horizontal='center')
            
    ws1.column_dimensions['A'].width = 3.5
    ws1.column_dimensions['B'].width = 7.0
    ws1.column_dimensions['C'].width = 34.0
    ws1.column_dimensions['D'].width = 3.0
    for d in range(1, days_cnt + 1):
        ws1.column_dimensions[get_column_letter(4 + d)].width = 4.2
    ws1.column_dimensions[tot_hrs_col].width = 14.0
    ws1.column_dimensions[work_col].width = 10.0
    
    # SHEET 2: Overtime
    ws2 = wb.create_sheet(title="Sheet2")
    ws2.views.sheetView[0].showGridLines = True
    
    start_c = 2
    for assoc in associates:
        emp_name = assoc["name"]
        days_vals = assoc.get("days", ["8"] * days_cnt)
        
        ws2.cell(row=6, column=start_c, value=emp_name).font = FONT_BOLD_10
        ws2.cell(row=7, column=start_c, value="OVER TIME - TIME STAMP").font = FONT_BOLD_9
        
        headers = ["DATE", "IN TIME", "OUT TIME", "NET HOURS", "EXCESS HOURS"]
        for h_i, h in enumerate(headers):
            c_cell = ws2.cell(row=8, column=start_c + h_i, value=h)
            c_cell.font = FONT_BOLD_9
            c_cell.border = THIN_BORDER
            c_cell.fill = HEADER_FILL
            c_cell.alignment = Alignment(horizontal='center')
            
        for d in range(1, days_cnt + 1):
            r = 8 + d
            raw_val = days_vals[d - 1] if d - 1 < len(days_vals) else "8"
            code = str(raw_val).strip().upper()
            is_work_8 = (code in ["8", 8])
            
            c_date = ws2.cell(row=r, column=start_c, value=d)
            c_in = ws2.cell(row=r, column=start_c + 1, value="08:00 AM" if is_work_8 else "-")
            c_out = ws2.cell(row=r, column=start_c + 2, value="05:00 PM" if is_work_8 else "-")
            c_net = ws2.cell(row=r, column=start_c + 3, value="08:00" if is_work_8 else "00:00")
            c_exc = ws2.cell(row=r, column=start_c + 4, value="00:00")
            
            for cell_obj in [c_date, c_in, c_out, c_net, c_exc]:
                cell_obj.border = THIN_BORDER
                cell_obj.alignment = Alignment(horizontal='center')
                cell_obj.font = FONT_REGULAR_9
                if not is_work_8 and code in color_map:
                    cell_obj.fill = color_map[code]
                    
        ws2.column_dimensions[get_column_letter(start_c)].width = 7.0
        ws2.column_dimensions[get_column_letter(start_c + 1)].width = 11.0
        ws2.column_dimensions[get_column_letter(start_c + 2)].width = 11.0
        ws2.column_dimensions[get_column_letter(start_c + 3)].width = 11.0
        ws2.column_dimensions[get_column_letter(start_c + 4)].width = 13.0
        
        start_c += 6
        
    wb.save(file_path)
    return file_path


def generate_master_excel(
    file_path: str,
    month_info: Dict[str, Any],
    associates: List[Dict[str, Any]],
    stamps: List[Dict[str, Any]] = None,
    companies_list: List[str] = None,
    properties_list: List[Dict[str, Any]] = None
):
    wb = Workbook()
    color_map = build_color_map(stamps or [])
    
    stamp_hours = {}
    if stamps:
        for st in stamps:
            c = str(st.get("code", "")).strip().upper()
            stamp_hours[c] = float(st.get("hours", 0))
    else:
        stamp_hours = {"8": 8, "WO": 8, "PH": 8, "SL": 8, "OFF": 8}
    paid_codes = [c for c, h in stamp_hours.items() if h > 0]
    
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.views.sheetView[0].showGridLines = True
    
    days_cnt = month_info["days_in_month"]
    
    ws1["A1"] = f"Monthly Attendance Sheet: {month_info['label'].upper()}"
    ws1["A1"].font = FONT_BOLD_14
    
    headers = [
        ("A", "Iqama No", 16.0),
        ("B", "ASSOCIATE NAME", 32.0),
        ("C", "Buyout company", 22.0),
        ("D", "Proparty", 12.0),
        ("E", "Department", 14.0),
        ("F", "Total_Hours", 14.0),
    ]
    for col_letter, text, width in headers:
        cell = ws1[f"{col_letter}2"]
        cell.value = text
        cell.font = FONT_BOLD_10
        cell.border = THIN_BORDER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws1.column_dimensions[col_letter].width = width
        
    for d in range(1, days_cnt + 1):
        col_l = get_column_letter(7 + d)
        cell = ws1[f"{col_l}2"]
        cell.value = d
        cell.font = FONT_BOLD_9
        cell.border = THIN_BORDER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws1.column_dimensions[col_l].width = 4.2
        
    current_row = 3
    first_d_col = get_column_letter(8)
    last_d_col = get_column_letter(7 + days_cnt)
    
    for assoc in associates:
        days_vals = assoc.get("days", ["8"] * days_cnt)
        
        ws1[f"A{current_row}"] = assoc.get("iqama", "")
        ws1[f"A{current_row}"].alignment = Alignment(horizontal='center')
        
        ws1[f"B{current_row}"] = assoc["name"]
        ws1[f"B{current_row}"].alignment = Alignment(horizontal='left')
        
        ws1[f"C{current_row}"] = assoc.get("company", "")
        ws1[f"C{current_row}"].alignment = Alignment(horizontal='left')
        
        ws1[f"D{current_row}"] = assoc.get("property", "RMH")
        ws1[f"D{current_row}"].alignment = Alignment(horizontal='center')
        
        ws1[f"E{current_row}"] = assoc.get("department", "0430")
        ws1[f"E{current_row}"].alignment = Alignment(horizontal='center')
        
        tot_parts = []
        if "8" in paid_codes:
            h = stamp_hours["8"]
            tot_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, 8)*{int(h) if h.is_integer() else h}")
            tot_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, \"8\")*{int(h) if h.is_integer() else h}")
        for pc in paid_codes:
            if pc != "8":
                h = stamp_hours[pc]
                tot_parts.append(f"COUNTIF({first_d_col}{current_row}:{last_d_col}{current_row}, \"{pc}\")*{int(h) if h.is_integer() else h}")
        ws1[f"F{current_row}"] = "=" + " + ".join(tot_parts) if tot_parts else "0"
        ws1[f"F{current_row}"].alignment = Alignment(horizontal='center')
        ws1[f"F{current_row}"].font = FONT_BOLD_10
        
        for base_c in ["A", "B", "C", "D", "E", "F"]:
            ws1[f"{base_c}{current_row}"].border = THIN_BORDER
            ws1[f"{base_c}{current_row}"].font = FONT_REGULAR_10 if base_c != "F" else FONT_BOLD_10
            
        for d_idx in range(days_cnt):
            raw_val = days_vals[d_idx] if d_idx < len(days_vals) else "8"
            code = str(raw_val).strip().upper()
            c_cell = ws1.cell(row=current_row, column=8 + d_idx)
            if code in ["8", 8]:
                c_cell.value = 8
                c_cell.font = FONT_REGULAR_9
            else:
                c_cell.value = code
                c_cell.font = FONT_BOLD_9
                if code in color_map:
                    c_cell.fill = color_map[code]
                    
            c_cell.alignment = Alignment(horizontal='center')
            c_cell.border = THIN_BORDER
                
        current_row += 1
        
    ws2 = wb.create_sheet(title="December")
    ws2["A1"] = "Company List"
    ws2["B1"] = "Proparty"
    ws2["C1"] = "Department"
    for c in ["A1", "B1", "C1"]:
        ws2[c].font = FONT_BOLD_10
        ws2[c].fill = HEADER_FILL
        ws2[c].border = THIN_BORDER
        
    comps = companies_list or ["Dar Al Ishraq", "Dar Al Wafa", "Nasim Alenjaz", "Sama Al riyadha", "Soundline", "Work Force", "Qasir AL Nokba", "Tamkeen"]
    for idx, comp in enumerate(comps, start=2):
        ws2[f"A{idx}"] = comp
        
    props = [p.get("code", "RMH") if isinstance(p, dict) else p for p in (properties_list or ["RMH", "MEA", "Housing"])]
    for idx, prop in enumerate(props, start=2):
        ws2[f"B{idx}"] = prop
        
    depts = ["Hotel Engineering", "Front Office", "Housekeeping", "Laundry"]
    for idx, d in enumerate(depts, start=2):
        ws2[f"C{idx}"] = d
        
    wb.save(file_path)
    return file_path
